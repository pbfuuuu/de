"""
МОДУЛЬ 1. Импорт данных заказчика из Excel (папка ../import) в базу SQLite.
Запуск: python import_data.py  ->  создаёт my_shop.db и скрипт БД my_shop_dump.sql
"""
import os
import sqlite3
from datetime import datetime

from openpyxl import load_workbook

DB_PATH = "my_shop.db"
# Данные (товары, цены, логины, пароли) в код НЕ пишем — они берутся из Excel в папке import.
IMPORT_DIR = "../import"


def get_or_create(cursor, table, name):
    """Возвращает id значения из справочника; если его нет — добавляет."""
    cursor.execute(f"SELECT id FROM {table} WHERE name = ?", (name,))
    row = cursor.fetchone()
    if row is not None:
        return row[0]
    cursor.execute(f"INSERT INTO {table} (name) VALUES (?)", (name,))
    return cursor.lastrowid


def to_date(value):
    """Дата из Excel -> текст ГГГГ-ММ-ДД. Неправильную дату печатаем — исправить её в Excel руками."""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    print("Неправильная дата:", value, "— исправь в Excel и запусти снова")
    return None


def clean(value):
    """Убирает пробелы и переносы строк по краям; None превращает в пустую строку."""
    return str(value).strip() if value is not None else ""


# ---------- создаём пустую базу по схеме ----------
if os.path.exists(DB_PATH):
    os.remove(DB_PATH)

conn = sqlite3.connect(DB_PATH)
conn.execute("PRAGMA foreign_keys = ON")       # включаем проверку внешних ключей
with open("schema.sql", encoding="utf-8") as f:
    conn.executescript(f.read())
cursor = conn.cursor()

# ---------- пользователи  >>> МЕНЯТЬ имя файла и порядок столбцов ----------
sheet = load_workbook(f"{IMPORT_DIR}/user_import.xlsx").active
for row in sheet.iter_rows(min_row=2, values_only=True):
    if row[0] is None:                          # пустые строки в конце файла
        continue
    role, full_name, login, password = row[:4]
    role_id = get_or_create(cursor, "role", clean(role))
    cursor.execute("INSERT INTO users (full_name, login, password, role_id) VALUES (?, ?, ?, ?)",
                   (clean(full_name), clean(login), clean(password), role_id))

# ---------- товары  >>> МЕНЯТЬ: порядок переменных = порядок столбцов Excel ----------
sheet = load_workbook(f"{IMPORT_DIR}/Tovar.xlsx").active
for row in sheet.iter_rows(min_row=2, values_only=True):
    if row[0] is None:
        continue
    article, name, unit, price, supplier, manufacturer, category, discount, quantity, description, photo = row[:11]
    category_id = get_or_create(cursor, "category", clean(category))
    manufacturer_id = get_or_create(cursor, "manufacturer", clean(manufacturer))
    supplier_id = get_or_create(cursor, "supplier", clean(supplier))
    cursor.execute(
        """INSERT INTO product (article, name, unit, description, photo, quantity, price, discount,
                                category_id, manufacturer_id, supplier_id)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
        (clean(article), clean(name), clean(unit), clean(description) or None, clean(photo) or None,
         quantity, price, discount, category_id, manufacturer_id, supplier_id))

# ---------- пункты выдачи: заголовка нет, номер пункта = номер строки  >>> МЕНЯТЬ имя файла ----------
sheet = load_workbook(f"{IMPORT_DIR}/Пункты выдачи_import.xlsx").active
for row in sheet.iter_rows(min_row=1, values_only=True):
    if row[0] is None:
        continue
    cursor.execute("INSERT INTO pickup_point (address) VALUES (?)", (clean(row[0]),))

# ---------- заказы и их состав  >>> МЕНЯТЬ порядок столбцов ----------
sheet = load_workbook(f"{IMPORT_DIR}/Заказ_import.xlsx").active
for row in sheet.iter_rows(min_row=2, values_only=True):
    if row[0] is None:
        continue
    number, items, order_date, delivery_date, point, client, code, status = row[:8]

    # клиента в заказе ищем по ФИО  >>> МЕНЯТЬ, если в файле логин или id
    cursor.execute("SELECT id FROM users WHERE full_name = ?", (clean(client),))
    user_id = cursor.fetchone()[0]
    status_id = get_or_create(cursor, "order_status", clean(status))

    cursor.execute(
        """INSERT INTO orders (id, order_date, delivery_date, pickup_code,
                               pickup_point_id, user_id, order_status_id)
           VALUES (?, ?, ?, ?, ?, ?, ?)""",
        (number, to_date(order_date), to_date(delivery_date), code, point, user_id, status_id))

    # "А112Т4, 2, G843H5, 2" -> пары (артикул, количество)  >>> МЕНЯТЬ, если состав заказа записан иначе
    parts = clean(items).split(",")
    for i in range(0, len(parts), 2):
        article = parts[i].strip()
        quantity = int(parts[i + 1])
        cursor.execute("SELECT id FROM product WHERE article = ?", (article,))
        product_id = cursor.fetchone()[0]
        cursor.execute("INSERT INTO order_item (order_id, product_id, quantity) VALUES (?, ?, ?)",
                       (number, product_id, quantity))

# ---------- сохраняем, проверяем, выгружаем скрипт БД ----------
conn.commit()
# >>> МЕНЯТЬ список таблиц под свою схему
for table in ("users", "role", "product", "category", "manufacturer", "supplier",
              "pickup_point", "orders", "order_status", "order_item"):
    print(table, "-", conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0])

with open("my_shop_dump.sql", "w", encoding="utf-8") as f:
    for line in conn.iterdump():
        f.write(line + "\n")
conn.close()
print("Готово: my_shop.db и my_shop_dump.sql")
